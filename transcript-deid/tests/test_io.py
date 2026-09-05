from pathlib import Path

from docx import Document as Docx

from transcript_deid import io
from transcript_deid.io.speakers import is_role_label, split_speaker

VTT = """WEBVTT

NOTE made by hand

1
00:00:01.000 --> 00:00:04.500 align:start
<v Interviewer>Hi Maria.</v>

00:00:04.600 --> 00:00:08.200
Maria Lopez: It's me.

00:00:08.300 --> 00:00:09.000
no speaker here
"""

SRT = """1
00:00:01,000 --> 00:00:04,500
Interviewer: Hi.

2
00:00:04,600 --> 00:00:08,200
Maria: Hello there.
"""


def test_split_speaker_variants():
    assert split_speaker("Interviewer: hi")[:2] == ("Interviewer", "hi")
    assert split_speaker("Maria Lopez (00:05): hi") == ("Maria Lopez", "hi", " (00:05):")
    assert split_speaker("Maria Lopez 00:05 hi") == ("Maria Lopez", "hi", " 00:05")
    assert split_speaker("[Maria]: hi")[0] == "Maria"
    assert split_speaker("Speaker 2: hi")[0] == "Speaker 2"
    assert split_speaker("Note: this is a long sentence that keeps going and going.")[0] == "Note"
    assert split_speaker("plain text with no label")[0] is None
    assert split_speaker("I went home at 5: it was late")[0] is None


def test_role_labels():
    assert is_role_label("Interviewer") and is_role_label("Speaker 2") and is_role_label("P")
    assert not is_role_label("Maria Lopez")


def test_vtt_round_trip(tmp_path: Path):
    src = tmp_path / "a.vtt"
    src.write_text(VTT, encoding="utf-8")
    doc = io.read(src)
    assert [s.speaker for s in doc.segments] == ["Interviewer", "Maria Lopez", None]
    assert doc.segments[0].meta["settings"] == " align:start"
    texts = [s.text for s in doc.segments]
    speakers = [s.speaker for s in doc.segments]
    out = io.write(doc, texts, speakers, tmp_path / "b.vtt")
    assert out.read_text(encoding="utf-8") == VTT
    # and a redacted write keeps the voice tag style
    speakers[1] = "FEAR-0102"
    io.write(doc, texts, speakers, tmp_path / "c.vtt")
    assert "FEAR-0102: It's me." in (tmp_path / "c.vtt").read_text()


def test_srt_round_trip(tmp_path: Path):
    src = tmp_path / "a.srt"
    src.write_text(SRT, encoding="utf-8")
    doc = io.read(src)
    assert doc.format == "srt" and doc.segments[1].speaker == "Maria"
    out = io.write(doc, [s.text for s in doc.segments], [s.speaker for s in doc.segments], tmp_path / "b.srt")
    assert out.read_text(encoding="utf-8") == SRT


def test_txt_round_trip(tmp_path: Path):
    src = tmp_path / "a.txt"
    src.write_text("Interviewer: hi\n\nMaria (00:01): hello\n", encoding="utf-8")
    doc = io.read(src)
    out = io.write(doc, [s.text for s in doc.segments], [s.speaker for s in doc.segments], tmp_path / "b.txt")
    assert out.read_text(encoding="utf-8") == "Interviewer: hi\n\nMaria (00:01): hello\n"


def test_docx_round_trip_preserves_styles(tmp_path: Path):
    src = tmp_path / "a.docx"
    d = Docx()
    d.add_heading("Interview", level=1)
    p = d.add_paragraph()
    p.add_run("Maria Lopez: ").bold = True
    p.add_run("I was born on 04/12/2008.")
    t = d.add_table(rows=1, cols=1)
    t.cell(0, 0).text = "Interviewer: table cell"
    d.save(src)

    doc = io.read(src)
    assert doc.format == "docx"
    assert doc.segments[1].speaker == "Maria Lopez"
    assert doc.segments[2].text == "table cell"
    texts = [s.text for s in doc.segments]
    speakers = [s.speaker for s in doc.segments]
    texts[1] = "I was born on [DATE]."
    speakers[1] = "FEAR-0102"
    out = io.write(doc, texts, speakers, tmp_path / "b.docx")
    d2 = Docx(str(out))
    assert d2.paragraphs[0].style.name.startswith("Heading")
    assert d2.paragraphs[1].text == "FEAR-0102: I was born on [DATE]."
    assert d2.paragraphs[1].runs[0].bold is True
    assert d2.tables[0].cell(0, 0).text == "Interviewer: table cell"


def test_unsupported_suffix(tmp_path: Path):
    p = tmp_path / "x.pdf"
    p.write_bytes(b"")
    try:
        io.read(p)
    except ValueError as e:
        assert "Unsupported" in str(e)
    else:
        raise AssertionError("expected ValueError")
