"""Build examples/interview_0102.docx from the plain-text lines (Otter-style layout)."""

from pathlib import Path

from docx import Document

lines = [
    ("Interviewer", "00:00", "Thanks for joining today, Maria."),
    ("Maria Lopez", "00:05", "Sure. My mom Elena said I should mention I was born on 04/12/2008."),
    ("Interviewer", "00:12", "Noted. Where do you spend most of your time?"),
    ("Maria Lopez", "00:15", "At Lincoln High School, or at my aunt's place in Round Rock."),
    ("Maria Lopez", "00:22", "Umm [unintelligible 00:24] the the the thing with Mr. Patel."),
]
d = Document()
d.add_heading("Interview FEAR-0102", level=1)
for spk, ts, text in lines:
    d.add_paragraph(f"{spk} ({ts}): {text}")
out = Path(__file__).with_name("interview_0102.docx")
d.save(out)
print("wrote", out)
