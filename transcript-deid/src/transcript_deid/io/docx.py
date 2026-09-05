"""Word (.docx) transcripts.

Every body paragraph (including those inside tables) becomes a ``Segment``.
On write we open the *original* file again and replace paragraph text in
place so styles, headers, footers and page setup survive.  Character-level
formatting inside a paragraph collapses to the first run's formatting when
the text changed; unchanged paragraphs are left untouched.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document as _Docx
from docx.text.paragraph import Paragraph

from ..models import Document, Segment
from .speakers import join_speaker, split_speaker


def _iter_paragraphs(d) -> list[Paragraph]:
    paras: list[Paragraph] = list(d.paragraphs)
    for table in d.tables:
        for row in table.rows:
            for cell in row.cells:
                paras.extend(cell.paragraphs)
    return paras


def read(path: Path) -> Document:
    d = _Docx(str(path))
    segments: list[Segment] = []
    for i, p in enumerate(_iter_paragraphs(d)):
        speaker, text, suffix = split_speaker(p.text)
        segments.append(
            Segment(
                index=i,
                text=text,
                speaker=speaker,
                meta={"style": p.style.name if p.style is not None else None, "original": p.text, "suffix": suffix},
            )
        )
    return Document(source=str(path), format="docx", segments=segments)


def _set_paragraph_text(p: Paragraph, text: str) -> None:
    runs = p.runs
    if not runs:
        p.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def write(doc: Document, texts: list[str], speakers: list[str | None], out: Path) -> None:
    d = _Docx(doc.source)
    paras = _iter_paragraphs(d)
    if len(paras) != len(doc.segments):
        raise RuntimeError(
            f"{doc.source}: paragraph count changed since it was read "
            f"({len(paras)} now vs {len(doc.segments)} recorded)"
        )
    for p, seg, text, speaker in zip(paras, doc.segments, texts, speakers):
        new = join_speaker(speaker, text, seg.meta.get("suffix", ":"))
        if new != seg.meta.get("original", p.text):
            _set_paragraph_text(p, new)
    d.save(str(out))
